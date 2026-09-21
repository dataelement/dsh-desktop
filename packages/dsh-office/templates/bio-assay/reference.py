import json,datetime,math
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment,Border,Side
from openpyxl.chart import LineChart,BarChart,ScatterChart,Reference,Series
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.chart.axis import ChartLines
from openpyxl.drawing.line import LineProperties
from openpyxl.drawing.text import Paragraph,ParagraphProperties,CharacterProperties,Font as DrawingFont
from openpyxl.chart.text import RichText
from openpyxl.worksheet.page import PageMargins
from openpyxl.formatting.rule import ColorScaleRule,CellIsRule
from openpyxl.utils import get_column_letter as col
from openpyxl.workbook.properties import CalcProperties
with open(office['inputs'][0],encoding='utf8') as f:data=json.load(f)
wb=Workbook();wb.remove(wb.active)
FONT='Noto Sans CJK SC';INK='253446';MUTED='687A8D';WHITE='FFFFFF';PAPER='F4F7FA';NAVY='174661';BLUE='286DAC';ACCENT='DE9C56';COLORS=[BLUE,'D99145','67A7AF','906EA7'];RED='B65759'
def cell(ws,r,c,value,fmt=None,bold=False,color=INK,size=10,fill=None):
 x=ws.cell(r,c,value);x.font=Font(name=FONT,size=size,color=color,bold=bold);x.alignment=Alignment(vertical='center',horizontal='left' if isinstance(value,str) and not value.startswith('=') else 'right',wrap_text=True)
 if fmt:x.number_format=fmt
 if fill:x.fill=PatternFill('solid',fgColor=fill)
 return x
def band(ws,r,c1,c2,value,size=11,bold=False,fill=None,color=INK):
 ws.merge_cells(start_row=r,start_column=c1,end_row=r,end_column=c2)
 if fill:
  for c in range(c1,c2+1):ws.cell(r,c).fill=PatternFill('solid',fgColor=fill)
 return cell(ws,r,c1,value,bold=bold,color=color,size=size,fill=fill)
def setup(name,rows=54,last='P',width=8.4,raw=False):
 ws=wb.create_sheet(name);ws.sheet_view.showGridLines=False;ws.sheet_view.zoomScale=85
 for c in range(1,19):ws.column_dimensions[col(c)].width=width
 if not raw:ws.column_dimensions['A'].width=2.5
 for r in range(1,rows+1):ws.row_dimensions[r].height=19
 ws.print_area=f'A1:{last}{rows}';ws.print_options.horizontalCentered=True;ws.sheet_properties.pageSetUpPr.fitToPage=True
 ws.page_setup.orientation='portrait';ws.page_setup.paperSize=ws.PAPERSIZE_A3;ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=1
 ws.page_margins=PageMargins(left=.28,right=.28,top=.25,bottom=.25,header=0,footer=.1)
 if raw:ws.freeze_panes='B8'
 return ws
def heading(ws,title,sub):
 band(ws,2,2,16,title,22,True,color=NAVY);ws.row_dimensions[2].height=37
 band(ws,3,2,16,sub,10,color=MUTED);ws.row_dimensions[3].height=25
 ws.sheet_properties.tabColor=NAVY
def section(ws,r,title):
 band(ws,r,2,16,title,11,True,fill=PAPER,color=NAVY);ws.row_dimensions[r].height=23

def kpi(ws,c,label,formula,fmt):
 band(ws,5,c,c+2,label,9,color=MUTED)
 x=band(ws,6,c,c+2,formula,23,True,color=NAVY);x.number_format=fmt;ws.row_dimensions[6].height=36

def rich(size=900):
 return RichText(p=[Paragraph(pPr=ParagraphProperties(defRPr=CharacterProperties(latin=DrawingFont(typeface=FONT),ea=DrawingFont(typeface=FONT),sz=size,solidFill=MUTED)),endParaRPr=CharacterProperties(lang='zh-CN'))])
def style(ch,title,width=12,height=7,legend=False):
 ch.title=title;ch.width=width;ch.height=height;ch.style=13;ch.graphical_properties=GraphicalProperties(solidFill=WHITE,ln=LineProperties(noFill=True))
 ch.x_axis.txPr=rich();ch.y_axis.txPr=rich();ch.x_axis.majorGridlines=None;ch.y_axis.majorGridlines=ChartLines(spPr=GraphicalProperties(ln=LineProperties(solidFill='E3EAF0',w=6000)))
 ch.title.txPr=rich(1100)
 if legend:ch.legend.position='b';ch.legend.txPr=rich(850)
 else:ch.legend=None
 for i,s in enumerate(ch.series):
  s.graphicalProperties.line.solidFill=COLORS[i%len(COLORS)];s.graphicalProperties.line.width=22000;s.graphicalProperties.solidFill=COLORS[i%len(COLORS)]
 return ch
def bar(ws,source,cols,r1,r2,category,title,anchor,width=12,height=7,stack=False):
 ch=BarChart()
 for c in cols:ch.add_data(Reference(source,min_col=c,min_row=r1-1,max_row=r2),titles_from_data=True)
 ch.set_categories(Reference(source,min_col=category,min_row=r1,max_row=r2));ch.y_axis.scaling.min=0;ch.gapWidth=70
 if stack:ch.grouping='stacked';ch.overlap=100
 style(ch,title,width,height,len(cols)>1);ws.add_chart(ch,anchor);return ch
def scatter(ws,source,xcol,ycols,r1,r2,title,anchor,width=24.4,height=8,labels=None,log=False):
 ch=ScatterChart();ch.scatterStyle='lineMarker'
 for n,c in enumerate(ycols):
  s=Series(Reference(source,min_col=c,min_row=r1,max_row=r2),Reference(source,min_col=xcol,min_row=r1,max_row=r2),title=labels[n] if labels else str(source.cell(r1-1,c).value));s.marker.symbol='circle' if r2-r1<40 else 'none';s.marker.size=4;ch.series.append(s)
 if log:ch.x_axis.scaling.logBase=10;ch.x_axis.scaling.min=.01;ch.x_axis.scaling.max=100
 style(ch,title,width,height,len(ycols)>1);ws.add_chart(ch,anchor);return ch

def headers(ws,r,labels):
 ws.row_dimensions[r].height=32
 for c,h in enumerate(labels,1):cell(ws,r,c,h,bold=True,color=WHITE,fill=NAVY,size=9).alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
def table_row(ws,r,vals,formats=None):
 for c,v in enumerate(vals,1):cell(ws,r,c,v,(formats or {}).get(c),fill=WHITE if r%2==0 else PAPER,size=10)
def save():
 wb.calculation=CalcProperties(calcMode='auto',fullCalcOnLoad=True);wb.save(office['output'])
NAVY='574065';BLUE='766298';ACCENT='3F9296';COLORS=['766298','3F9296','D59A60','AB7FA3'];PAPER='F7F5F9'
O=setup('实验总览');D=setup('剂量响应',52,'H',15,True);Q=setup('板内质控',57,'M',9.5,True);R=setup('原始读数',46,'G',17,True)
heading(O,data['title'],'荧光酶活分析    3 个独立批次 · 每剂量每批次 8 个技术重复    模拟数据')
for ws,title,sub,lastcol in [(D,'剂量响应与批次统计','批次均值先分别归一化，再计算跨批次均值和样本 SD；n = 3',8),(Q,'板内质控与孔位分布','每板 96 孔；第 1 列空白，第 2 列溶剂对照，第 3–12 列为十档浓度',13),(R,'原始荧光读数','288 个模拟孔位全部保留；此预览展示前 39 条原始读数',7)]:
 band(ws,2,1,lastcol,title,20,True,color=NAVY);band(ws,3,1,lastcol,sub,10,color=MUTED);ws.row_dimensions[2].height=34;ws.row_dimensions[3].height=30
headers(R,7,data['fields'])
for r,w in enumerate(data['wells'],8):table_row(R,r,w,{4:'0.00',5:'0',6:'#,##0.0'})
R.column_dimensions['A'].width=11;R.column_dimensions['B'].width=12;R.column_dimensions['C'].width=17;R.column_dimensions['G'].width=38;R.auto_filter.ref='A7:G295'
band(R,4,1,7,'模拟方法：固定种子 41865；浓度响应叠加批次差异和测量噪声，记录中保留一处偏差。',9,color=MUTED);R.row_dimensions[4].height=30
headers(Q,7,['批次','空白均值','对照均值','对照 SD','对照 CV','空白 SD','信号窗口'])
for p in range(3):
 r=8+p;start=8+p*96
 vals=[f'P{p+1}',f'=AVERAGE(\'原始读数\'!F{start}:F{start+7})',f'=AVERAGE(\'原始读数\'!F{start+8}:F{start+15})',f'=STDEV(\'原始读数\'!F{start+8}:F{start+15})',f'=D{r}/C{r}',f'=STDEV(\'原始读数\'!F{start}:F{start+7})',f'=C{r}-B{r}']
 table_row(Q,r,vals,{2:'0.0',3:'#,##0',4:'0.0',5:'0.0%',6:'0.0',7:'#,##0'})
cell(Q,8,9,'技术重复 CV 阈值',size=9);cell(Q,8,12,data['cvThreshold'],'0%',color=BLUE);Q.column_dimensions['I'].width=16
band(Q,11,1,13,'阈值用于演示复核标记；高 CV 孔组保留参与计算，并回查原始记录。',9,color=MUTED)
for p in range(3):
 top=13+p*14;band(Q,top,1,13,f'P{p+1}  背景扣除后的相对活性 / %',11,True,fill=PAPER,color=NAVY)
 for c in range(1,13):cell(Q,top+1,c+1,c,'0',bold=True,color=MUTED)
 for row in range(8):
  cell(Q,top+2+row,1,chr(65+row),bold=True,color=MUTED)
  for c in range(12):
   raw=8+p*96+c*8+row;cell(Q,top+2+row,c+2,f'=(\'原始读数\'!F{raw}-$B${8+p})/$G${8+p}','0%')
 Q.conditional_formatting.add(f'B{top+2}:M{top+9}',ColorScaleRule(start_type='num',start_value=0,start_color='F3EEF7',mid_type='num',mid_value=.5,mid_color='B1A0C5',end_type='num',end_value=1,end_color='527F97'))
 band(Q,top+11,2,13,'孔位颜色按原始读数和本板对照实时计算；所有测量值均保留。',9,color=MUTED)
Q.row_dimensions[57].height=10
headers(D,7,['浓度 / µM','P1 活性','P2 活性','P3 活性','批次均值','批次 SD','平均抑制率','最大组内 CV'])
for i,dose in enumerate(data['doses']):
 r=8+i;cell(D,r,1,dose,'General',color=BLUE)
 cv=[]
 for p in range(3):
  raw=8+p*96+16+i*8;qr=8+p
  cell(D,r,2+p,f'=(AVERAGE(\'原始读数\'!F{raw}:F{raw+7})-\'板内质控\'!B{qr})/\'板内质控\'!G{qr}','0.0%')
  cv.append(f'STDEV(\'原始读数\'!F{raw}:F{raw+7})/(AVERAGE(\'原始读数\'!F{raw}:F{raw+7})-\'板内质控\'!B{qr})')
 for c,expr in {5:f'=AVERAGE(B{r}:D{r})',6:f'=STDEV(B{r}:D{r})',7:f'=1-E{r}',8:'=MAX('+','.join(cv)+')'}.items():cell(D,r,c,expr,'0.0%',fill=PAPER if i%2 else WHITE)
D.conditional_formatting.add('H8:H17',CellIsRule(operator='greaterThan',formula=["'板内质控'!$L$8"],fill=PatternFill('solid',fgColor='F4DFD8'),font=Font(color=RED)))
band(D,20,1,8,'统计口径',11,True,color=NAVY)
band(D,22,1,8,'相对活性 =（样品均值 − 本板空白均值）÷（本板溶剂对照均值 − 本板空白均值）。',10);D.row_dimensions[22].height=33
band(D,24,1,8,'跨批次离散程度使用 3 个独立批次均值的样本标准差。技术重复仅用于评价板内测量稳定性。',10);D.row_dimensions[24].height=33
band(D,26,1,8,'曲线连接实测浓度点，横轴使用对数刻度。50% 抑制浓度仅报告相邻剂量区间。',10);D.row_dimensions[26].height=33
band(D,28,1,8,'方法参考：NIH Assay Guidance Manual / Basic Guidelines for Reporting Non-Clinical Data',9,color=MUTED)
band(D,29,1,8,data['methodUrl'],8,color=MUTED)
ch=bar(D,D,[6],8,17,1,'跨批次样本 SD / 相对活性','A33',24,9);ch.y_axis.numFmt='0%'
kpi(O,2,'孔位读数 / 个',"=COUNTA('原始读数'!B8:B295)",'0');kpi(O,6,'独立批次 / 个',"=COUNTA('板内质控'!A8:A10)",'0');kpi(O,10,'每剂量技术重复 / 批次',"=COUNT('原始读数'!F24:F31)",'0');kpi(O,14,'剂量水平 / 档',"=COUNT('剂量响应'!A8:A17)",'0')
section(O,9,'01  剂量升高时，相对酶活性如何变化')
ch=scatter(O,D,1,[2,3,4,5],8,17,'浓度响应 / µM 与相对活性','B10',24.4,8.3,['批次 P1','批次 P2','批次 P3','跨批次均值'],True);ch.y_axis.numFmt='0%';ch.y_axis.scaling.min=0;ch.y_axis.scaling.max=1.15
section(O,28,'02  对照稳定性与组内重复差异')
ch=bar(O,Q,[5],8,10,1,'溶剂对照 CV / %','B29',11.8,7);ch.y_axis.numFmt='0%'
ch=bar(O,D,[8],8,17,1,'各剂量最大组内 CV / %','J29',11.8,7);ch.y_axis.numFmt='0%'
section(O,44,'关键读数与复核事项')
for c,h in zip([2,6,10,14],['浓度 / µM','平均活性','批次 SD','复核标记']):band(O,45,c,min(c+2,16),h,9,True,fill=NAVY,color=WHITE)
for r,src in enumerate([8,13,14,17],46):
 for c,expr,fmt in [(2,f"='剂量响应'!A{src}",'0.00'),(6,f"='剂量响应'!E{src}",'0.0%'),(10,f"='剂量响应'!F{src}",'0.0%'),(14,f'=IF(\'剂量响应\'!H{src}>\'板内质控\'!$L$8,"复核 CV","正常")','@')]:band(O,r,c,min(c+2,16),expr,10,fill=PAPER if r%2 else WHITE).number_format=fmt
band(O,51,2,16,'50% 抑制区间 / µM',10,True,color=NAVY)
# Locate the first observed point that reaches 50%; the discrete interval updates with inputs.
# Check monotonicity before reporting a discrete crossing interval.
band(O,52,2,16,'=IF(SUMPRODUCT(--(\'剂量响应\'!G9:G17<\'剂量响应\'!G8:G16))>0,"响应有回升，逐段复核",IF(COUNTIFS(\'剂量响应\'!G8:G17,"<0.5")=0,"低于最低测试浓度",IF(COUNTIFS(\'剂量响应\'!G8:G17,"<0.5")=10,"高于最高测试浓度",INDEX(\'剂量响应\'!A8:A17,COUNTIFS(\'剂量响应\'!G8:G17,"<0.5"))&" 至 "&INDEX(\'剂量响应\'!A8:A17,COUNTIFS(\'剂量响应\'!G8:G17,"<0.5")+1))))',12,True,color=NAVY)
band(O,54,2,16,'模拟读数用于展示分析结构；偏差记录与原始孔位保留，可在后续工作表核对。',9,color=MUTED)
save()
