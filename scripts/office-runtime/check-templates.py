"""Independent native-file checks for the generated example catalogue."""
import argparse
import hashlib
import json
import math
import struct
import zipfile
from pathlib import Path
from openpyxl import load_workbook
from docx import Document
from PIL import Image

FONT_BUDGET = 10 * 1024 * 1024
BANNED_FONT_NAMES = ('FangS-SC', 'ShuS-SC', 'HYKaiTiJ', 'HYZhongJianHeiJ', '方正', '汉仪')
BANNED_CONTENT = ('中国中车', '吉安委办联', '质量安全月', '关于开展2025')
PRIMARY_NAME_IDS = {1, 2, 3, 4, 6, 16, 17, 18, 21, 22}
RESERVED_FONT_NAMES = ('Source Han', 'SourceHan', '思源', 'Noto')
APPROVED_NOTICE_FONT = {
    'familyName': 'DSH Gov Song',
    'postScriptName': 'DSHGovSong',
    'styleName': 'Regular',
    'outlineWeight': 600,
    'subsetSha256': 'dd90fea8c670db387e9014dda49e0bcb504553d431a968f4f1d521815b640f64',
    'subsetBytes': 3139120,
    'sourceSha256': '04f163fb935acd9c0caa79e535a43d510a724d25e584d36212e494a9f9db64f4',
    'licenseSha256': '9ff5bb567e1b92c801fc1069e5fbf992ff8efccacb9db94e5959a5b3ba9bb903',
}
W_NS = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main',
        'r': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'}


def deobfuscate_font(data, font_key):
    # ECMA-376 reverses the fontKey GUID bytes and XORs them across the first 32 bytes.
    raw = bytes.fromhex(font_key.strip().strip('{}').replace('-', ''))
    if len(raw) != 16:
        raise ValueError('embedded font key is not a GUID')
    key = raw[::-1]
    prefix = bytearray(data[:32])
    for index in range(len(prefix)):
        prefix[index] ^= key[index % 16]
    return bytes(prefix) + data[32:]


def font_tables(data):
    if len(data) < 12:
        raise ValueError('embedded font is truncated')
    scaler, count = struct.unpack('>IH', data[:6])
    tables = {}
    for index in range(count):
        start = 12 + index * 16
        tag, _checksum, offset, length = struct.unpack('>4sIII', data[start:start + 16])
        tables[tag.decode('latin1')] = (offset, length)
    return scaler, tables


def font_name_records(data, tables):
    offset, length = tables['name']
    table = data[offset:offset + length]
    count, string_offset = struct.unpack('>HH', table[2:6])
    records = []
    for index in range(count):
        start = 6 + index * 12
        platform, encoding, _language, name_id, size, rel = struct.unpack('>HHHHHH', table[start:start + 12])
        raw = table[string_offset + rel:string_offset + rel + size]
        if platform == 3 and encoding in (1, 10):
            text = raw.decode('utf-16-be')
        elif platform == 0:
            text = raw.decode('utf-16-be')
        elif platform == 1 and encoding == 0:
            text = raw.decode('mac-roman', errors='replace')
        else:
            continue
        records.append((name_id, text))
    return records


def audit_docx(path, source_dir, template_id=None):
    """Reject commercial embeds, unknown subsets and the retired enterprise sample."""
    source_dir = Path(source_dir)
    notice = template_id == 'government-notice'
    embedded = 0
    family = None
    if notice:
        for file in source_dir.iterdir():
            if not file.is_file() or file.name in {'font-source.json', 'OFL.txt'} or file.suffix.lower() not in {'.md', '.json', '.txt'}:
                continue
            text = file.read_text(encoding='utf-8')
            for banned in BANNED_FONT_NAMES + BANNED_CONTENT:
                if banned in text:
                    raise AssertionError((file.name, banned))
    with zipfile.ZipFile(path) as archive:
        names = archive.namelist()
        if notice and any(name.startswith('word/media/') for name in names):
            raise AssertionError((path.name, 'government notice must not carry media'))
        xml_parts = [name for name in names if name.endswith('.xml')]
        for name in xml_parts:
            text = archive.read(name).decode('utf-8', errors='replace')
            banned_values = BANNED_FONT_NAMES + (BANNED_CONTENT if notice else ())
            for banned in banned_values:
                if banned in text:
                    raise AssertionError((path.name, name, banned))
        font_parts = [name for name in names if name.startswith('word/fonts/')]
        if not font_parts:
            if notice:
                raise AssertionError((path.name, 'government notice is missing its approved font'))
            return {'embedded_fonts': 0, 'family': None}
        manifest_path = source_dir / 'font-source.json'
        if not manifest_path.is_file():
            raise AssertionError((path.name, 'embedded font has no approved manifest'))
        manifest = json.loads(manifest_path.read_text(encoding='utf-8'))
        approved_manifest = {
            'familyName': manifest.get('familyName'),
            'postScriptName': manifest.get('postScriptName'),
            'styleName': manifest.get('styleName'),
            'outlineWeight': manifest.get('outlineWeight'),
            'subsetSha256': manifest.get('subsetSha256'),
            'subsetBytes': manifest.get('subsetBytes'),
            'sourceSha256': manifest.get('source', {}).get('sha256'),
        }
        expected_manifest = {key: value for key, value in APPROVED_NOTICE_FONT.items() if key != 'licenseSha256'}
        if approved_manifest != expected_manifest:
            raise AssertionError((path.name, 'manifest does not match the approved notice font'))
        license_data = (source_dir / 'OFL.txt').read_bytes()
        license_text = license_data.decode('utf-8')
        if 'SIL Open Font License' not in license_text or 'Version 1.1' not in license_text:
            raise AssertionError((path.name, 'OFL 1.1 attribution is missing'))
        if hashlib.sha256(license_data).hexdigest() != APPROVED_NOTICE_FONT['licenseSha256']:
            raise AssertionError((path.name, 'OFL license does not match the approved notice font'))
        font_table = archive.read('word/fontTable.xml')
        from lxml import etree
        root = etree.fromstring(font_table)
        embeds = root.xpath('//w:embedRegular', namespaces=W_NS)
        if len(embeds) != len(font_parts) or len(embeds) != 1:
            raise AssertionError((path.name, 'only one approved embedded font is allowed'))
        embed = embeds[0]
        if embed.get('{%s}subsetted' % W_NS['w']) not in ('true', '1', 'on'):
            raise AssertionError((path.name, 'embedded font must be marked subsetted'))
        declared = embed.getparent().get('{%s}name' % W_NS['w'])
        if declared != manifest['familyName']:
            raise AssertionError((path.name, declared, manifest['familyName']))
        relationship = embed.get('{%s}id' % W_NS['r'])
        rels = etree.fromstring(archive.read('word/_rels/fontTable.xml.rels'))
        target = rels.xpath('//*[local-name()="Relationship" and @Id=$id]/@Target', id=relationship)
        if len(target) != 1:
            raise AssertionError((path.name, 'embedded font relationship is missing'))
        font_name = 'word/' + target[0].lstrip('/') if not target[0].startswith('word/') else target[0]
        payload = deobfuscate_font(archive.read(font_name), embed.get('{%s}fontKey' % W_NS['w']))
        if len(payload) > FONT_BUDGET or len(payload) != APPROVED_NOTICE_FONT['subsetBytes'] \
                or hashlib.sha256(payload).hexdigest() != APPROVED_NOTICE_FONT['subsetSha256']:
            raise AssertionError((path.name, 'embedded font does not match the approved subset'))
        scaler, tables = font_tables(payload)
        if scaler != 0x00010000 or 'glyf' not in tables or 'fvar' in tables or 'CFF ' in tables:
            raise AssertionError((path.name, 'embedded font must be a static TrueType glyf font'))
        records = font_name_records(payload, tables)
        families = {text for name_id, text in records if name_id == 1}
        postscript_names = {text for name_id, text in records if name_id == 6}
        styles = {text for name_id, text in records if name_id in (2, 17)}
        if families != {APPROVED_NOTICE_FONT['familyName']} \
                or postscript_names != {APPROVED_NOTICE_FONT['postScriptName']} \
                or styles != {APPROVED_NOTICE_FONT['styleName']}:
            raise AssertionError((path.name, families))
        for name_id, text in records:
            if name_id in PRIMARY_NAME_IDS and any(reserved in text for reserved in RESERVED_FONT_NAMES):
                raise AssertionError((path.name, 'modified font still uses a reserved name', name_id, text))
        embedded = len(payload)
        family = APPROVED_NOTICE_FONT['familyName']
    return {'embedded_fonts': embedded, 'family': family}

def check(root, sources):
    catalog = json.loads((root / 'catalog.json').read_text(encoding='utf-8'))
    assert {entry.name for entry in root.iterdir()} == {'catalog.json', *(item['id'] for item in catalog)}, 'catalogue resource root'
    numeric_checks, previews, visual_checks = 0, 0, 0
    for item in catalog:
        directory = root / item['id']
        for name, digest in item['files'].items():
            assert hashlib.sha256((directory / name).read_bytes()).hexdigest() == digest, (item['id'], name)
        previews += item['pages']
        sample = directory / item['sampleFile']
        with zipfile.ZipFile(sample) as archive:
            assert archive.testzip() is None
            assert all('\\' not in name and not name.startswith('/') for name in archive.namelist())
            if item['mode'] == 'excel':
                assert all(entry.date_time == (2026, 9, 1, 0, 0, 0) for entry in archive.infolist()), item['id']
                assert b'2026-09-01T00:00:00Z' in archive.read('docProps/core.xml'), item['id']
        if item['mode'] == 'word':
            audit_docx(sample, directory, item['id'])
            document = Document(sample)
            assert document.paragraphs, item['id']
            proof = directory / 'preview-checks.json'
            if proof.exists():
                checks = json.loads(proof.read_text(encoding='utf-8'))
                assert checks['schemaVersion'] == 1 and isinstance(checks['checks'], list)
                for row in checks['checks']:
                    assert row['image'] in item['files'] and row['image'].startswith('pages/')
                    x0, y0, x1, y1 = row['region']
                    assert 0 <= x0 < x1 <= 1 and 0 <= y0 < y1 <= 1
                    assert 0 <= row['darkMax'] <= 255 and 0 < row['minInkCoverage'] <= 1
                    with Image.open(directory / row['image']) as source:
                        image = source.convert('RGB')
                    region = image.crop((int(image.width*x0), int(image.height*y0), int(image.width*x1), int(image.height*y1)))
                    dark = sum(1 for pixel in region.getdata() if max(pixel) < row['darkMax'])
                    coverage = dark / (region.width * region.height)
                    assert coverage >= row['minInkCoverage'], (item['id'], row['label'], coverage,
                        'Generate and verify previews with complete native font access')
                    visual_checks += 1
            continue
        values, editable = load_workbook(sample, data_only=True), load_workbook(sample, data_only=False)
        for sheet in values:
            for row in sheet:
                for cell in row:
                    assert cell.data_type != 'e', (item['id'], sheet.title, cell.coordinate, cell.value)
        assert sum(len(sheet._charts) for sheet in editable) > 0, item['id']
        legacy = sources / 'legacy' / item['id'] / item['sampleFile']
        if legacy.exists():
            original = load_workbook(legacy, data_only=False)
            assert editable.sheetnames == original.sheetnames
            for sheet in editable:
                assert sheet.freeze_panes == original[sheet.title].freeze_panes
                assert len(sheet._charts) == len(original[sheet.title]._charts)
        else:
            assert all(sheet.freeze_panes == 'A2' for sheet in editable)
        if (directory / 'checks.json').exists():
            for row in json.loads((directory / 'checks.json').read_text(encoding='utf-8')):
                actual = values[row['sheet']][row['cell']].value
                assert math.isclose(actual, row['expected'], abs_tol=row.get('tolerance', 1e-6)), (item['id'], row, actual)
                numeric_checks += 1
    return {'templates': len(catalog), 'preview_pages': previews, 'numeric_checks': numeric_checks, 'preview_visual_checks': visual_checks,
            'resource_hashes': 'PASS', 'native_zip_crc': 'PASS', 'excel_error_cells': 0,
            'native_charts': 'PASS', 'frozen_panes': 'PASS'}


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('root', type=Path)
    parser.add_argument('sources', type=Path)
    args = parser.parse_args()
    print(json.dumps(check(args.root, args.sources)))
